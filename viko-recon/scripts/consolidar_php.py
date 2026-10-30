#!/usr/bin/env python3
"""Merge v1.6 report.json files offline. Does not read inventories or Vault."""
import argparse
import collections
import datetime
import json
import os
from pathlib import Path
import report

def merge(paths):
    selected = {}
    audit_source = None
    for path in paths:
        data = json.loads(Path(path).read_text(encoding='utf-8-sig'))
        if data.get('schema_version') != 2 or 'hosts' not in data or 'findings' not in data:
            raise ValueError('Se requiere reporte.json del recolector 1.6: ' + str(path))
        for host in data['hosts']:
            alias = host['host']
            stamp = report.report_datetime(host.get('timestamp_utc')) or report.report_datetime(data.get('generated_utc'))
            previous = selected.get(alias)
            if previous is None or stamp > previous[0]:
                selected[alias] = (stamp, host, [f for f in data['findings'] if f['host'] == alias])
        if audit_source is None or (data.get('access_records_count') or 0) > (audit_source.get('access_records_count') or 0):
            audit_source = data
    hosts = [selected[k][1] for k in sorted(selected)]
    findings = [f for k in sorted(selected) for f in selected[k][2]]
    result = {'schema_version': 2, 'generated_utc': datetime.datetime.now(datetime.timezone.utc).isoformat(),
              'rules_reviewed_on': report.RULESET['reviewed_on'], 'checks': ['php'],
              'scope_note': report.RULESET['scope_note'], 'hosts': hosts, 'findings': findings,
              'inventory_audit': (audit_source or {}).get('inventory_audit', []),
              'access_records_count': (audit_source or {}).get('access_records_count'),
              'summary': {'hosts': len(hosts), 'hosts_not_evaluated': sum(h['collection_status'] != 'recolectado' for h in hosts),
                  'hosts_partial': sum(report.incomplete_coverage(h['coverage'], ['php']) for h in hosts),
                  'hosts_requiring_review': sum(h['review_required'] for h in hosts),
                  'cve_observation_rows': len(findings), 'statuses': dict(collections.Counter(f['status'] for f in findings))}}
    result = report.enrich_php(result)
    for row in result['inventory_audit']:
        row['selected_this_job'] = row['host'] in selected
    return result

def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--report', action='append', required=True, help='Repetir para cada reporte.json de los pilotos')
    p.add_argument('--out', required=True)
    p.add_argument('--xlsx', action='store_true', help='Requiere openpyxl instalado; no instala paquetes')
    a=p.parse_args(); os.umask(0o077)
    result=merge(a.report)
    report.write_reports(result,a.out)
    if a.xlsx:
        try: import openpyxl
        except ImportError: p.error('CSV generados. XLSX requiere openpyxl en el equipo donde consolide.')
        import csv
        from openpyxl.styles import Font, PatternFill, Alignment
        wb=openpyxl.Workbook(); wb.remove(wb.active)
        for name, csvname in [('Equipos','resumen.csv'),('Componentes_CVE','hallazgos.csv'),('Cruce_inventario','cruce_inventario.csv')]:
            ws=wb.create_sheet(name)
            with open(Path(a.out)/csvname,encoding='utf-8-sig',newline='') as stream:
                for row in csv.reader(stream): ws.append(row)
            ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
            for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='07556B')
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width=min(65,max(16,max(len(str(c.value or '')) for c in col[:30])+2))
            for row in ws.iter_rows(min_row=2):
                for c in row: c.alignment=Alignment(vertical='top',wrap_text=True)
        wb.save(Path(a.out)/'Consolidado_PHP_VIKO.xlsx')
    print('Consolidado: {} equipos; {} filas CVE. Archivos en {}'.format(len(result['hosts']),len(result['findings']),a.out))
if __name__=='__main__': main()
